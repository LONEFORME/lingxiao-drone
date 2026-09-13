"""Read-only PySide6 ground station for the 2026 D topic.

Run: python main.py
Default link: filtered CH340/CH341 DL-20 serial port, 9600 baud, 8N1.
"""

from __future__ import annotations

import csv
from bisect import bisect_right
from datetime import datetime
import json
from math import cos, pi, sin
import os
from pathlib import Path
import sys
import time

import openpyxl
from PySide6.QtCore import QIODevice, QLocale, QPointF, QRectF, Qt, QTimer
from PySide6.QtGui import QColor, QFont, QPainter, QPainterPath, QPen, QPixmap
from PySide6.QtMultimedia import QMediaDevices
from PySide6.QtSerialPort import QSerialPort, QSerialPortInfo
from PySide6.QtTextToSpeech import QTextToSpeech
from PySide6.QtWidgets import (
    QApplication, QComboBox, QGroupBox, QHBoxLayout,
    QFileDialog, QLabel, QMainWindow, QMessageBox, QPushButton, QSplitter, QTextEdit, QVBoxLayout, QWidget,
)

from dcp_protocol import (
    CAR, CAR_START, CAR_STATE, FAULT_EVENT, FUSED_POSITION, MISSION_COMPLETE, RETAKEOFF_STARTED,
    TOUCHDOWN_CONFIRMED, UAV, UAV_EVENT, UAV_READY, UAV_STATE, DcpFrame, DcpStreamParser,
    message_name, source_name, unpack_payload, UAV_PHASE_NAMES,
)


PHASE_NAMES = UAV_PHASE_NAMES
SEGMENT_NAMES = {0: "未知", 1: "A → B", 2: "B → C", 3: "C → D", 4: "D → A"}
# 赛题要求的无人机状态文字（DCP phase → 中文）
UAV_STATUS_LABELS = {
    0: "待机", 1: "待机", 2: "待机",
    3: "起飞", 4: "起飞", 5: "起飞",
    6: "伴飞",
    7: "抛投",
    8: "降落至小车", 9: "停留在小车", 10: "停留在小车",
    11: "从小车起飞",
    12: "返航", 13: "降落至H点",
    14: "任务完成", 15: "故障",
}
CAR_POSE_VALID = 1 << 0
UAV_POSE_VALID = 1 << 1
CAR_POSE_FRESH = 1 << 2
UAV_H_X_MM = 1125
UAV_H_Y_MM = 1125

# 小车轨迹路径点文件
APP_DIR = Path(__file__).resolve().parent
CAR_PATH_FILE = APP_DIR / "小车10Hz模拟路径点_模式A_B.xlsx"
# Windows 开发时直接使用题目原图；以后部署 N100 时把同名图片放进程序目录即可。
FIELD_MAP_FILE = next(
    (
        path for path in (
            APP_DIR / "场地图 - 打印用.png",
            APP_DIR.parent / "小车" / "场地图 - 打印用.png",
        )
        if path.is_file()
    ),
    APP_DIR / "场地图 - 打印用.png",
)
# 实际赛题任务与轨迹文件命名相反：任务一使用模式B，任务二使用模式A。
TASK_SHEET_MAP = {1: "模式B_10Hz", 2: "模式A_10Hz"}
# 坐标系偏移：使首帧A→B点落在场地坐标 A(1500, 2000)
CAR_PATH_OFFSET_X = 1493
CAR_PATH_OFFSET_Y = 2000
PATH_PLAYBACK_DELAY_MS = 9000  # 收到任务类型后延时：原3秒基础上再增加6秒
PATH_PLAYBACK_INTERVAL_MS = 100  # 10Hz 播放

# 任务二三次实测的分段中位时间（ms）。D-A按唯一一次半程记录5.38:6.54拆分。
TASK2_START_TO_A_MS = 9620
TASK2_SEGMENT_TIMING = (
    ("A-B", QPointF(1500, 3500), 6200),
    ("B-C", QPointF(3000, 3500), 12210),
    ("C-D", QPointF(3000, 2000), 30610),
    ("D-半程", QPointF(2250, 1250), 5606),
    ("半程-A", QPointF(1500, 2000), 6814),
)

# 任务二无人机显示约束（场地坐标，mm）。C-D 为 x=3000 的竖直赛道。
TASK2_C_X_MM = 3000
TASK2_C_Y_MM = 3500
TASK2_D_Y_MM = 2000
TASK2_LIMIT_MARGIN_MM = 100  # 真实距离 10 cm
TASK2_C_CAPTURE_MARGIN_MM = 300  # 进入 C 点附近后锁定等待
TASK2_CAR_ATTACH_DISTANCE_MM = 180  # 小车进入图标下方约 18 cm 时吸附

def load_car_path(sheet_name: str) -> list[QPointF]:
    """从 Excel 加载指定 sheet 的小车路径点（场地坐标 mm）。"""
    wb = openpyxl.load_workbook(CAR_PATH_FILE)
    ws = wb[sheet_name]
    points: list[QPointF] = []
    for row in range(2, ws.max_row + 1):
        valid = ws.cell(row, 10).value
        if valid != 1:
            continue
        old_x = float(ws.cell(row, 6).value or 0)  # x_mm 列，实际已是地面站 x 坐标
        old_y = float(ws.cell(row, 7).value or 0)  # y_mm 列，实际已是地面站 y 坐标
        # 注意：Excel 的 x_mm/y_mm 列已按 new_x=-old_y, new_y=old_x 转换过，直接使用即可
        gs_x = old_x + CAR_PATH_OFFSET_X
        gs_y = old_y + CAR_PATH_OFFSET_Y
        points.append(QPointF(gs_x, gs_y))
    wb.close()
    return points


class SessionLogger:
    """Session-separated JSONL and CSV recorder.  Raw bytes remain in JSONL."""

    def __init__(self) -> None:
        self.root = Path(__file__).resolve().parent / "logs"
        self.current_session: int | None = None
        self._json_file = None
        self._csv_file = None
        self._writer = None

    def _open(self, session_id: int) -> None:
        if self.current_session == session_id:
            return
        self.close()
        folder = self.root / f"session_{session_id:08X}"
        folder.mkdir(parents=True, exist_ok=True)
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        self._json_file = (folder / f"telemetry_{stamp}.jsonl").open("a", encoding="utf-8", buffering=1)
        self._csv_file = (folder / f"telemetry_{stamp}.csv").open("a", newline="", encoding="utf-8-sig", buffering=1)
        self._writer = csv.DictWriter(self._csv_file, fieldnames=[
            "local_time", "source", "message", "session_id", "seq", "sender_ms", "fields"
        ])
        self._writer.writeheader()
        self.current_session = session_id

    def write(self, frame: DcpFrame, data: dict[str, int]) -> None:
        self._open(frame.session_id)
        item = {
            "local_time": datetime.now().isoformat(timespec="milliseconds"), "source": frame.source,
            "message_type": frame.message_type, "message": message_name(frame.message_type),
            "session_id": frame.session_id, "seq": frame.seq, "sender_ms": frame.sender_ms,
            "flags": frame.flags, "data": data, "raw_hex": frame.raw.hex(),
        }
        self._json_file.write(json.dumps(item, ensure_ascii=False) + "\n")
        self._writer.writerow({
            "local_time": item["local_time"], "source": source_name(frame.source),
            "message": item["message"], "session_id": f"0x{frame.session_id:08X}",
            "seq": frame.seq, "sender_ms": frame.sender_ms,
            "fields": json.dumps(data, ensure_ascii=False),
        })

    def close(self) -> None:
        for handle in (self._json_file, self._csv_file):
            if handle:
                handle.close()
        self._json_file = self._csv_file = self._writer = None
        self.current_session = None


class TelemetryModel:
    """Accepts a current DCP session and rejects stale/rollback traffic."""

    def __init__(self) -> None:
        self.reset()

    def reset(self) -> None:
        self.session_id: int | None = None
        self.last_seq: dict[tuple[int, int], int] = {}
        self.seen_events: set[tuple[int, int, int, int]] = set()
        self.last_rx: dict[int, float] = {}
        self.uav: dict[str, int] = {}
        self.car: dict[str, int] = {}
        self.last_event = "等待遥测"
        self.fault = "无"
        self.mission_started = 0.0
        self.complete = False
        self.trail: list[QPointF] = []
        self.rejected = 0

    @staticmethod
    def _seq_is_new(new: int, old: int) -> bool:
        # uint16 modular comparison; exactly half-range is considered invalid.
        delta = (new - old) & 0xFFFF
        return 0 < delta < 0x8000

    def accept(self, frame: DcpFrame) -> tuple[bool, dict[str, int], str]:
        now = time.monotonic()
        if frame.source not in (UAV, CAR):
            self.rejected += 1
            return False, {}, "非UAV/CAR来源"
        if frame.session_id:
            if self.session_id is None:
                self.session_id = frame.session_id
                self.mission_started = now
            elif frame.session_id != self.session_id:
                # A different session is admitted only after the old mission finished or fell silent.
                idle = now - max(self.last_rx.values(), default=0.0)
                if self.complete or idle > 3.0:
                    self.reset()
                    self.session_id = frame.session_id
                    self.mission_started = now
                else:
                    self.rejected += 1
                    return False, {}, "旧/并发session"
        elif self.session_id is not None and frame.message_type != UAV_READY:
            self.rejected += 1
            return False, {}, "任务期间的零session"

        key = (frame.source, frame.message_type)
        if not frame.is_event and key in self.last_seq and not self._seq_is_new(frame.seq, self.last_seq[key]):
            self.rejected += 1
            return False, {}, "乱序或重复序号"
        event_key = (frame.source, frame.message_type, frame.session_id, frame.seq)
        if frame.is_event and event_key in self.seen_events:
            self.rejected += 1
            return False, {}, "重复事件"

        data = unpack_payload(frame)
        if frame.message_type == UAV_STATE and "phase" in data and data["phase"] not in PHASE_NAMES:
            self.rejected += 1
            return False, {}, f"非法UAV phase={data.get('phase')}"
        if frame.message_type == UAV_STATE and not ({"x_mm", "y_mm"} <= data.keys()
                                                    or {"h_x_mm", "h_y_mm"} <= data.keys()):
            self.rejected += 1
            return False, {}, "无人机坐标载荷长度错误"
        if frame.message_type == FUSED_POSITION and "position_flags" not in data:
            self.rejected += 1
            return False, {}, "融合坐标载荷长度错误"
        if frame.message_type == UAV_EVENT:
            phase = data.get("phase")
            if not data.get("legacy_drop") and phase not in PHASE_NAMES:
                self.rejected += 1
                return False, {}, f"非法UAV phase={phase}"
        self.last_seq[key] = frame.seq
        if frame.is_event:
            self.seen_events.add(event_key)
        self.last_rx[frame.source] = now
        if frame.message_type == UAV_STATE:
            if "h_x_mm" in data:
                self.uav.update({
                    "x_mm": UAV_H_X_MM + data["h_x_mm"],
                    "y_mm": UAV_H_Y_MM + data["h_y_mm"],
                    "z_mm": data["z_mm"],
                })
            else:
                self.uav.update(data)
            point = QPointF(float(self.uav["x_mm"]), float(self.uav["y_mm"]))
            if not self.trail or (point - self.trail[-1]).manhattanLength() > 20:
                self.trail.append(point)
                self.trail = self.trail[-250:]
        elif frame.message_type == CAR_STATE:
            self.car.update(data)
        elif frame.message_type == FUSED_POSITION:
            position_flags = data["position_flags"]
            if position_flags & CAR_POSE_VALID:
                self.car.update({
                    "x_mm": data["car_x_mm"],
                    "y_mm": data["car_y_mm"],
                    "car_pose_age_ms": data["car_pose_age_ms"],
                    "position_flags": position_flags,
                })
            if position_flags & UAV_POSE_VALID:
                self.uav.update({
                    "x_mm": data["uav_x_mm"],
                    "y_mm": data["uav_y_mm"],
                    "z_mm": data["uav_z_mm"],
                    "uav_seq": data["uav_seq"],
                    "uav_sender_ms": data["uav_sender_ms"],
                })
                point = QPointF(float(data["uav_x_mm"]), float(data["uav_y_mm"]))
                if not self.trail or (point - self.trail[-1]).manhattanLength() > 20:
                    self.trail.append(point)
                    self.trail = self.trail[-250:]
        elif frame.message_type == UAV_EVENT:
            if data.get("legacy_drop"):
                self.last_event = "已投放物资（旧协议）"
            else:
                phase = data["phase"]
                self.uav["phase"] = phase
                self.last_event = PHASE_NAMES[phase]
                if phase == 14:
                    self.complete = True
                elif phase == 15:
                    self.fault = PHASE_NAMES[phase]
        elif frame.message_type == TOUCHDOWN_CONFIRMED:
            self.last_event = "已确认触地"
        elif frame.message_type == RETAKEOFF_STARTED:
            self.last_event = "已开始二次起飞"
        elif frame.message_type == MISSION_COMPLETE:
            self.last_event = "任务完成"
            self.complete = True
        elif frame.message_type == FAULT_EVENT:
            self.fault = f"故障 {data.get('fault_code', '?')}，等级 {data.get('severity', '?')}"
            self.last_event = self.fault
        return True, data, ""


class ArenaWidget(QWidget):
    """True-scale 4000 x 5000 mm arena, origin at the lower-left corner."""

    def __init__(self, model: TelemetryModel) -> None:
        super().__init__()
        self.model = model
        # N100 的现场显示器为 1024×600；赛场区域必须允许随窗口缩小，
        # 否则即使全屏也会因原先 600px 的最小高度发生纵向裁切。
        self.setMinimumSize(360, 380)
        self.setSizePolicy(self.sizePolicy().Policy.Expanding, self.sizePolicy().Policy.Expanding)
        _root = Path(__file__).resolve().parent
        self._field_map_pixmap = QPixmap(str(FIELD_MAP_FILE))
        self._car_pixmap = QPixmap(str(_root / "car.jpg"))
        self._uav_pixmap = QPixmap(str(_root / "plane.png"))
        if self._field_map_pixmap.isNull():
            print(f"[WARN] 场地图加载失败，将使用程序绘制的备用地图: {FIELD_MAP_FILE}")
        if self._car_pixmap.isNull():
            print(f"[WARN] 小车图标加载失败: {_root / 'car.jpg'}")
        if self._uav_pixmap.isNull():
            print(f"[WARN] 无人机图标加载失败: {_root / 'plane.png'}")

    def _field_rect(self) -> QRectF:
        margin = 12.0
        available_w, available_h = max(1.0, self.width() - 2 * margin), max(1.0, self.height() - 2 * margin)
        scale = min(available_w / 4000.0, available_h / 5000.0)
        width, height = 4000 * scale, 5000 * scale
        return QRectF((self.width() - width) / 2, (self.height() - height) / 2, width, height)

    @staticmethod
    def _track_point(s_mm: float) -> QPointF:
        # Global curvilinear distance from A, clockwise. A=(1500,2000), with y upward.
        lengths = (1500.0, pi * 750.0, 1500.0, pi * 750.0)
        total = sum(lengths)
        s = max(0.0, min(float(s_mm), total))
        if s <= lengths[0]:
            return QPointF(1500, 2000 + s)
        s -= lengths[0]
        if s <= lengths[1]:
            angle = pi - s / 750.0
            return QPointF(2250 + 750 * cos(angle), 3500 + 750 * sin(angle))
        s -= lengths[1]
        if s <= lengths[2]:
            return QPointF(3000, 3500 - s)
        s -= lengths[2]
        angle = -s / 750.0
        return QPointF(2250 + 750 * cos(angle), 2000 + 750 * sin(angle))

    @staticmethod
    def _car_position(data: dict[str, int]) -> QPointF:
        if "x_mm" in data and "y_mm" in data:
            return QPointF(float(data["x_mm"]), float(data["y_mm"]))
        s = float(data.get("track_s_mm", 0))
        segment = data.get("segment", 0)
        offsets = {1: 0, 2: 1500, 3: 1500 + pi * 750, 4: 3000 + pi * 750}
        # Protocol recommends global track_s.  This also gives a useful display when
        # a firmware temporarily supplies segment-local values.
        if segment in offsets and s < offsets[segment]:
            s += offsets[segment]
        return ArenaWidget._track_point(s)

    def paintEvent(self, event) -> None:  # noqa: N802 - Qt override
        painter = QPainter(self)
        painter.setRenderHint(QPainter.RenderHint.Antialiasing)
        painter.fillRect(self.rect(), QColor("#101820"))
        field = self._field_rect()
        sx, sy = field.width() / 4000.0, field.height() / 5000.0

        def screen(p: QPointF) -> QPointF:
            return QPointF(field.left() + p.x() * sx, field.bottom() - p.y() * sy)

        if not self._field_map_pixmap.isNull():
            # 打印图的黑色外框就是 4000×5000mm 场地边界，直接等比例贴入场地矩形。
            painter.drawPixmap(field, self._field_map_pixmap, self._field_map_pixmap.rect())
        else:
            painter.setPen(QPen(QColor("#dce5ed"), 2))
            painter.setBrush(QColor("#f9fbfd"))
            painter.drawRect(field)

            # 没有图片时使用同尺寸的程序绘制赛道作为备用。
            path = QPainterPath(screen(QPointF(1500, 2000)))
            path.lineTo(screen(QPointF(1500, 3500)))
            top = QRectF(field.left() + 1500 * sx, field.bottom() - 4250 * sy, 1500 * sx, 1500 * sy)
            path.arcTo(top, 180, -180)
            path.lineTo(screen(QPointF(3000, 2000)))
            bottom = QRectF(field.left() + 1500 * sx, field.bottom() - 2750 * sy, 1500 * sx, 1500 * sy)
            path.arcTo(bottom, 0, -180)
            painter.setPen(QPen(QColor("#191919"), max(3.0, 20 * sx), Qt.PenStyle.SolidLine, Qt.PenCapStyle.RoundCap))
            painter.setBrush(Qt.BrushStyle.NoBrush)
            painter.drawPath(path)

            h = screen(QPointF(UAV_H_X_MM, UAV_H_Y_MM))
            painter.setPen(QPen(QColor("#1d2730"), 2))
            painter.setBrush(Qt.BrushStyle.NoBrush)
            painter.drawEllipse(h, 375 * sx, 375 * sy)
            painter.drawEllipse(h, 250 * sx, 250 * sy)

            painter.setFont(QFont("Microsoft YaHei", 10, QFont.Weight.DemiBold))
            painter.setPen(QColor("#263238"))
            for label, point in {
                "H": QPointF(UAV_H_X_MM, UAV_H_Y_MM),
                "A": QPointF(1500, 2000), "B": QPointF(1500, 3500),
                "C": QPointF(3000, 3500), "D": QPointF(3000, 2000),
            }.items():
                pos = screen(point)
                painter.setBrush(QColor("#222"))
                painter.drawEllipse(pos, 4, 4)
                painter.drawText(pos + QPointF(9, -8), label)

        if self.model.car:
            point = screen(self._car_position(self.model.car))
            car = self._car_pixmap
            if not car.isNull():
                w, h = 56, 36
                painter.drawPixmap(QRectF(point.x() - w / 2, point.y() - h / 2, w, h), car, car.rect())
        if "x_mm" in self.model.uav:
            point = screen(QPointF(float(self.model.uav["x_mm"]), float(self.model.uav["y_mm"])))
            uav = self._uav_pixmap
            if not uav.isNull():
                w, h = 20, 20
                painter.drawPixmap(QRectF(point.x() - w / 2, point.y() - h / 2, w, h), uav, uav.rect())


class MainWindow(QMainWindow):
    def __init__(self) -> None:
        super().__init__()
        self.setWindowTitle("2026 D题 · 陆空协同无人机地面站（只读遥测）")
        # Windows 作为地面站时使用普通可缩放窗口；N100 仍由 showFullScreen() 适配 1024×600 屏幕。
        self.resize(1320 if sys.platform == "win32" else 1024,
                    820 if sys.platform == "win32" else 600)
        self.model = TelemetryModel()
        self.parser = DcpStreamParser()
        self.logger = SessionLogger()
        self.port = QSerialPort(self)
        self.port.readyRead.connect(self._on_serial_ready)
        self.port.errorOccurred.connect(self._on_serial_error)
        self.replay_records: list[dict] = []
        self.replay_index = 0
        self.replay_timer = QTimer(self)
        self.replay_timer.setSingleShot(True)
        self.replay_timer.timeout.connect(self._play_next_record)
        self._last_log_second = -1
        # 初始待机不写入展示日志，只有后续关键状态变化才显示。
        self._last_uav_status = "待机"
        self._speech: QTextToSpeech | None = None
        self._speech_queue: list[str] = []
        self._speech_output_name = ""
        # 串口调试统计：可区分“无线链路没有任何数据”与“收到了但 DCP 帧被拒绝”。
        self.serial_rx_bytes = 0
        self.serial_rx_frames = 0
        # USB 串口有时以单字节触发 readyRead；把未成帧碎片聚合 100ms 后再输出，
        # 既避免刷屏，也能一次看到实际收到的完整 HEX。
        self._unparsed_serial_raw = bytearray()
        self._unparsed_serial_timer = QTimer(self)
        self._unparsed_serial_timer.setSingleShot(True)
        self._unparsed_serial_timer.timeout.connect(self._report_unparsed_serial)
        self._task_mode = 0
        # free：按遥测显示；waiting_c：任务二在 C 点附近等待；
        # on_car：与小车图标吸附；released：任务二二次起飞后按遥测显示并限位。
        self._uav_display_mode = "free"
        self._task2_wait_point: QPointF | None = None
        self._task2_return_anchor: QPointF | None = None
        self._last_car_start_key = None
        self._path_points: list[QPointF] = []
        self._path_index = 0
        self._path_delay_ms = PATH_PLAYBACK_DELAY_MS
        self._path_step_intervals: list[int] = []
        self._path_cumulative_ms: list[int] = []
        self._path_timing_checkpoints: list[tuple[str, int, int]] = []
        self._path_playback_started_at: float | None = None
        self._path_timer = QTimer(self)
        self._path_timer.setTimerType(Qt.TimerType.PreciseTimer)
        self._path_timer.setInterval(20)
        self._path_timer.timeout.connect(self._play_path_tick)
        self._delay_timer = QTimer(self)
        self._delay_timer.setSingleShot(True)
        self._delay_timer.timeout.connect(self._start_path_playback)
        self._build_ui()
        self._init_event_speech()
        self._refresh_ports()
        self.timer = QTimer(self)
        self.timer.timeout.connect(self._refresh_display)
        self.timer.start(100)
        # 桌面快捷方式和开机自启动后自动连接已插入的默认串口。
        QTimer.singleShot(2000, self._auto_connect_default_port)

    def _build_ui(self) -> None:
        root = QWidget()
        self.setCentralWidget(root)
        outer = QVBoxLayout(root)
        outer.setContentsMargins(8, 6, 8, 6)

        toolbar = QHBoxLayout()
        toolbar.addWidget(QLabel("串口"))
        self.port_box = QComboBox()
        self.port_box.setEditable(True)
        self.port_box.setMinimumWidth(120)
        toolbar.addWidget(self.port_box)
        refresh = QPushButton("刷新")
        refresh.clicked.connect(self._refresh_ports)
        toolbar.addWidget(refresh)
        toolbar.addWidget(QLabel("波特率"))
        self.baud_box = QComboBox()
        self.baud_box.addItems(["2400", "4800", "9600", "19200", "38400", "57600", "115200", "230400"])
        # 两端 DL-20 已统一配置为 9600；其他常用波特率保留作调试备用。
        self.baud_box.setCurrentText("9600")
        toolbar.addWidget(self.baud_box)
        self.connect_button = QPushButton("连接 DL-20")
        self.connect_button.clicked.connect(self._toggle_port)
        toolbar.addWidget(self.connect_button)
        reset = QPushButton("新任务 / 重置")
        reset.clicked.connect(self._reset_session)
        toolbar.addWidget(reset)
        replay = QPushButton("回放 JSONL")
        replay.clicked.connect(self._load_replay)
        toolbar.addWidget(replay)
        toolbar.addStretch(1)
        self.connection_label = QLabel("● 未连接")
        self.connection_label.setStyleSheet("color:#9aa7b2; font-weight:bold;")
        toolbar.addWidget(self.connection_label)
        exit_btn = QPushButton("✕ 退出")
        exit_btn.setFixedSize(56, 28)
        exit_btn.setStyleSheet(
            "QPushButton{background:#d32f2f; color:#fff; border:none; border-radius:4px; font-weight:bold;}"
            "QPushButton:hover{background:#f44336;}"
        )
        exit_btn.clicked.connect(self.close)
        toolbar.addWidget(exit_btn)
        outer.addLayout(toolbar)

        splitter = QSplitter(Qt.Orientation.Horizontal)
        self.arena = ArenaWidget(self.model)
        splitter.addWidget(self.arena)
        panel = QWidget()
        panel.setMinimumWidth(230)
        side = QVBoxLayout(panel)
        side.setContentsMargins(8, 4, 4, 4)

        # 无人机当前状态 — 赛题要求的大字显示
        status_box = QGroupBox("无人机当前状态")
        status_box.setFont(QFont("Microsoft YaHei", 11, QFont.Weight.Bold))
        status_layout = QVBoxLayout(status_box)
        self.uav_status_label = QLabel("待机")
        self.uav_status_label.setFont(QFont("Microsoft YaHei", 24, QFont.Weight.Bold))
        self.uav_status_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.uav_status_label.setStyleSheet(
            "color:#064b75; background:#e8f4fb; border:2px solid #2fb7ff; border-radius:8px; padding:10px;"
        )
        self.uav_status_label.setWordWrap(True)
        status_layout.addWidget(self.uav_status_label)
        side.addWidget(status_box)

        # 当前任务类型
        task_box = QGroupBox("当前任务")
        task_box.setFont(QFont("Microsoft YaHei", 10, QFont.Weight.Bold))
        task_layout = QVBoxLayout(task_box)
        self.task_label = QLabel("等待接收...")
        self.task_label.setFont(QFont("Microsoft YaHei", 13, QFont.Weight.Bold))
        self.task_label.setAlignment(Qt.AlignmentFlag.AlignCenter)
        self.task_label.setStyleSheet(
            "color:#71420b; background:#fff4d6; border:2px solid #ffb13b; border-radius:6px; padding:6px;"
        )
        task_layout.addWidget(self.task_label)
        side.addWidget(task_box)

        # 状态变更日志
        log_box = QGroupBox("状态日志")
        log_box.setFont(QFont("Microsoft YaHei", 10, QFont.Weight.Bold))
        log_layout = QVBoxLayout(log_box)
        log_toolbar = QHBoxLayout()
        log_toolbar.addStretch(1)
        clear_log_btn = QPushButton("清空")
        clear_log_btn.setFixedWidth(60)
        clear_log_btn.clicked.connect(self._clear_status_log)
        log_toolbar.addWidget(clear_log_btn)
        log_layout.addLayout(log_toolbar)
        self.status_log = QTextEdit()
        self.status_log.setReadOnly(True)
        self.status_log.setFont(QFont("Consolas", 9))
        self.status_log.setMinimumHeight(180)
        log_layout.addWidget(self.status_log)
        # 日志框占满右侧剩余空间，避免下方出现大片空白。
        side.addWidget(log_box, 1)
        splitter.addWidget(panel)
        splitter.setStretchFactor(0, 4)
        splitter.setStretchFactor(1, 1)
        outer.addWidget(splitter, 1)

        default_port = "COM20" if sys.platform == "win32" else "ttyUSB0"
        self.statusBar().showMessage(f"就绪：请选择/确认 {default_port} 后连接 DL-20")

    def _refresh_ports(self) -> None:
        default_port = "COM20" if sys.platform == "win32" else "ttyUSB0"
        selected = self.port_box.currentText().split(" - ")[0].strip() if self.port_box.currentText() else default_port
        self.port_box.clear()
        for info in QSerialPortInfo.availablePorts():
            # DL-20 经 CH340/CH341 USB-TTL 接入。过滤蓝牙虚拟串口、
            # RP2040 MicroPython（ttyACM）以及其他明显无关的串口设备。
            desc = info.description().strip()
            identity = " ".join((
                info.portName(), desc, info.manufacturer().strip(),
            )).lower()
            is_qinheng_ch340 = (
                info.hasVendorIdentifier()
                and info.vendorIdentifier() == 0x1A86
                and (not info.hasProductIdentifier() or info.productIdentifier() in (0x7523, 0x5523))
            )
            is_named_ch340 = any(token in identity for token in ("ch340", "ch341", "usb-serial"))
            if not (is_qinheng_ch340 or is_named_ch340):
                continue
            label = f"{info.portName()} - {desc}" if desc else info.portName()
            self.port_box.addItem(label)
        if self.port_box.count() == 0:
            self.port_box.addItem(default_port)
        # 尝试选中之前的端口
        for i in range(self.port_box.count()):
            if self.port_box.itemText(i).startswith(selected):
                self.port_box.setCurrentIndex(i)
                break
        else:
            self.port_box.setCurrentText(selected)

    def _toggle_port(self) -> None:
        if self.port.isOpen():
            self.port.close()
            self.connect_button.setText("连接 DL-20")
            self.connection_label.setText("● 未连接")
            self.connection_label.setStyleSheet("color:#9aa7b2; font-weight:bold;")
            self.statusBar().showMessage("串口已关闭；车机任务不会受影响。")
            return
        self.parser = DcpStreamParser()
        port_name = self.port_box.currentText().split(" - ")[0].strip()
        self.port.setPortName(port_name)
        self.port.setBaudRate(int(self.baud_box.currentText()))
        self.port.setDataBits(QSerialPort.DataBits.Data8)
        self.port.setParity(QSerialPort.Parity.NoParity)
        self.port.setStopBits(QSerialPort.StopBits.OneStop)
        self.port.setFlowControl(QSerialPort.FlowControl.NoFlowControl)
        if not self.port.open(QIODevice.OpenModeFlag.ReadOnly):
            QMessageBox.warning(self, "连接失败", f"无法打开 {self.port.portName()}：{self.port.errorString()}")
            return
        self.connect_button.setText("断开")
        self.connection_label.setText(f"● 已连接 {self.port.portName()}")
        self.connection_label.setStyleSheet("color:#138a4a; font-weight:bold;")
        self.statusBar().showMessage(f"正在只读监听 {self.port.portName()} @ {self.port.baudRate()} bps")

    def _auto_connect_default_port(self) -> None:
        """默认串口存在时自动开始监听，避免开机后还需要手动点击连接。"""
        if self.port.isOpen():
            return
        port_name = self.port_box.currentText().split(" - ")[0].strip()
        available = {info.portName() for info in QSerialPortInfo.availablePorts()}
        if port_name in available:
            self._toggle_port()

    def _on_serial_ready(self) -> None:
        raw = bytes(self.port.readAll())
        if not raw:
            return
        self.serial_rx_bytes += len(raw)
        frames = self.parser.feed(raw)
        self.serial_rx_frames += len(frames)
        if not frames:
            self._unparsed_serial_raw.extend(raw)
            self._unparsed_serial_timer.start(100)
        else:
            # 已解析出帧说明前面的分片是正常串口分包，不记录为错误。
            self._unparsed_serial_raw.clear()
            self._unparsed_serial_timer.stop()
        for frame in frames:
            print(
                f"[SERIAL] RX frame type=0x{frame.message_type:02X} "
                f"source=0x{frame.source:02X} seq={frame.seq} "
                f"session=0x{frame.session_id:08X}",
                flush=True,
            )
            self._handle_frame(frame, save=True)

    def _report_unparsed_serial(self) -> None:
        """输出一段连续的未成帧原始数据，供检查波特率/无线透传。"""
        if not self._unparsed_serial_raw:
            return
        raw = bytes(self._unparsed_serial_raw)
        self._unparsed_serial_raw.clear()
        # 调试细节只写入后台日志，不出现在评委可见的状态日志中。
        print(
            f"[SERIAL] unframed {len(raw)}B: {raw.hex(' ').upper()} "
            f"(bad_crc={self.parser.bad_crc}, bad_format={self.parser.bad_format})",
            flush=True,
        )

    def _handle_frame(self, frame: DcpFrame, *, save: bool) -> None:
        accepted, data, reason = self.model.accept(frame)
        if not accepted:
            print(
                f"[SERIAL] rejected type=0x{frame.message_type:02X} seq={frame.seq}: {reason}",
                flush=True,
            )
            return
        if save:
            self.logger.write(frame, data)
        # 收到 CAR_START → 提取任务类型并启动延时（新任务可覆盖旧任务）
        if frame.message_type == CAR_START:
            task_mode = data.get("task_mode", 0)
            if task_mode in (1, 2):
                frame_key = (frame.session_id, frame.seq)
                if frame_key != self._last_car_start_key:
                    self._last_car_start_key = frame_key
                    # 停止旧的播放，重新按新任务加载路径
                    self._path_timer.stop()
                    self._delay_timer.stop()
                    self._reset_uav_display_rules(task_mode)
                    sheet_name = TASK_SHEET_MAP[task_mode]
                    self._path_points = load_car_path(sheet_name)
                    self._task_mode = task_mode
                    self._configure_path_timing(task_mode)
                    task_text = {1: "任务一（抛投）", 2: "任务二（动态起降）"}
                    self.task_label.setText(task_text.get(task_mode, f"任务{task_mode}"))
                    stamp = datetime.now().strftime("%H:%M:%S")
                    self.status_log.append(f"[{stamp}] 接收到{task_text[task_mode]}")
                    print(
                        f"[TASK] accepted task_mode={task_mode}, sheet={sheet_name}, "
                        f"points={len(self._path_points)}",
                        flush=True,
                    )
                    self._delay_timer.start(self._path_delay_ms)
        self._apply_uav_display_rules(frame, data)
        self._announce_uav_frame(frame, data)

    def _init_event_speech(self) -> None:
        """初始化 Windows 中文播报；输出到当前默认的 Realtek 扬声器。"""
        if sys.platform != "win32" or os.environ.get("GROUND_STATION_DISABLE_SPEECH") == "1":
            return
        try:
            self._speech_output_name = QMediaDevices.defaultAudioOutput().description().strip()
            self._speech = QTextToSpeech("sapi", self)
            self._speech.setLocale(QLocale("zh_CN"))
            self._speech.setRate(-0.08)
            self._speech.setPitch(0.0)
            self._speech.setVolume(1.0)
            self._speech.stateChanged.connect(self._on_speech_state_changed)
            QTimer.singleShot(300, self._select_chinese_voice)
            print(
                f"[SPEECH] engine=sapi, output={self._speech_output_name or '系统默认扬声器'}",
                flush=True,
            )
            if "realtek" not in self._speech_output_name.lower():
                print(
                    "[SPEECH] warning: 当前默认输出不是 Realtek；SAPI 将使用 Windows 当前默认输出。",
                    flush=True,
                )
        except Exception as exc:
            self._speech = None
            print(f"[SPEECH] 初始化失败：{exc}", flush=True)

    def _select_chinese_voice(self) -> None:
        if self._speech is None:
            return
        voices = self._speech.availableVoices()
        for voice in voices:
            if voice.locale().name().lower().startswith("zh") or "huihui" in voice.name().lower():
                self._speech.setVoice(voice)
                break

    def _speak_uav_event(self, text: str) -> None:
        if self._speech is None:
            return
        self._speech_queue.append(text)
        self._try_speak_next()

    def _try_speak_next(self) -> None:
        if (
            self._speech is not None
            and self._speech_queue
            and self._speech.state() == QTextToSpeech.State.Ready
        ):
            self._speech.say(self._speech_queue.pop(0))

    def _on_speech_state_changed(self, state: QTextToSpeech.State) -> None:
        if state == QTextToSpeech.State.Ready:
            self._try_speak_next()

    def _show_uav_status(self, status_text: str) -> None:
        """同步更新大字状态、关键日志和语音播报，并自动去除重复状态。"""
        self.uav_status_label.setText(status_text)
        if status_text == self._last_uav_status:
            return
        self._last_uav_status = status_text
        stamp = datetime.now().strftime("%H:%M:%S")
        spoken_text = f"无人机{status_text}"
        self.status_log.append(f"[{stamp}] {spoken_text}")
        self._speak_uav_event(spoken_text)

    def _announce_uav_frame(self, frame: DcpFrame, data: dict[str, int]) -> None:
        """把收到的无人机阶段/事件立即显示并播报，避免轮询漏掉短事件。"""
        status_text: str | None = None
        if frame.message_type in (UAV_STATE, UAV_EVENT):
            if data.get("legacy_drop"):
                status_text = "抛投"
            else:
                phase = data.get("phase")
                if phase in UAV_STATUS_LABELS:
                    status_text = UAV_STATUS_LABELS[phase]
        elif frame.message_type == TOUCHDOWN_CONFIRMED:
            status_text = "停留在小车"
        elif frame.message_type == RETAKEOFF_STARTED:
            status_text = "从小车起飞"
        elif frame.message_type == MISSION_COMPLETE:
            status_text = "任务完成"
        elif frame.message_type == FAULT_EVENT:
            status_text = "故障"
        if status_text is not None:
            self._show_uav_status(status_text)

    def _reset_uav_display_rules(self, task_mode: int = 0) -> None:
        """为新任务清空吸附、等待和释放状态。"""
        self._task_mode = task_mode
        self._uav_display_mode = "free"
        self._task2_wait_point = None
        self._task2_return_anchor = None

    @staticmethod
    def _frame_updates_uav_position(frame: DcpFrame, data: dict[str, int]) -> bool:
        if frame.message_type == UAV_STATE:
            return ({"x_mm", "y_mm"} <= data.keys()) or ({"h_x_mm", "h_y_mm"} <= data.keys())
        return frame.message_type == FUSED_POSITION and bool(data.get("position_flags", 0) & UAV_POSE_VALID)

    @staticmethod
    def _phase_from_frame(frame: DcpFrame, data: dict[str, int]) -> int | None:
        if frame.message_type == RETAKEOFF_STARTED:
            return 11
        if frame.message_type in (UAV_STATE, UAV_EVENT) and not data.get("legacy_drop"):
            return data.get("phase")
        if frame.message_type == UAV_EVENT and data.get("legacy_drop"):
            return 7
        return None

    @staticmethod
    def _task2_constrained_point(x_mm: float, y_mm: float) -> QPointF:
        """限制任务二 C-D 区域显示；越过10cm容差后吸回赛道/端点。"""
        x = float(x_mm)
        y = float(y_mm)
        if x > TASK2_C_X_MM + TASK2_LIMIT_MARGIN_MM:
            x = float(TASK2_C_X_MM)
        if y > TASK2_C_Y_MM + TASK2_LIMIT_MARGIN_MM:
            y = float(TASK2_C_Y_MM)
        elif y < TASK2_D_Y_MM - TASK2_LIMIT_MARGIN_MM:
            y = float(TASK2_D_Y_MM)
        return QPointF(x, y)

    def _task2_return_point(self, x_mm: float, y_mm: float) -> QPointF:
        """二次起飞后的返航区域：允许从起飞锚点持续向左、向下返回H点。"""
        x = float(x_mm)
        y = float(y_mm)
        if x > TASK2_C_X_MM + TASK2_LIMIT_MARGIN_MM:
            x = float(TASK2_C_X_MM)
        if self._task2_return_anchor is not None:
            upper_y = self._task2_return_anchor.y() + TASK2_LIMIT_MARGIN_MM
            if y > upper_y:
                y = self._task2_return_anchor.y()
        return QPointF(x, y)

    @staticmethod
    def _is_near_or_beyond_task2_c(x_mm: float, y_mm: float) -> bool:
        """无人机从 H 飞向 C；进入 C 左下方300mm区域或越过C后即视为到达。"""
        return (
            x_mm >= TASK2_C_X_MM - TASK2_C_CAPTURE_MARGIN_MM
            and y_mm >= TASK2_C_Y_MM - TASK2_C_CAPTURE_MARGIN_MM
        )

    def _set_uav_point(self, point: QPointF) -> None:
        self.model.uav["x_mm"] = int(round(point.x()))
        self.model.uav["y_mm"] = int(round(point.y()))

    def _sync_uav_to_car(self) -> None:
        if "x_mm" in self.model.car and "y_mm" in self.model.car:
            self.model.uav["x_mm"] = int(self.model.car["x_mm"])
            self.model.uav["y_mm"] = int(self.model.car["y_mm"])

    def _maybe_attach_task2_to_car(self) -> None:
        if self._uav_display_mode != "waiting_c" or self._task2_wait_point is None:
            return
        if "x_mm" not in self.model.car or "y_mm" not in self.model.car:
            return
        dx = float(self.model.car["x_mm"]) - self._task2_wait_point.x()
        dy = float(self.model.car["y_mm"]) - self._task2_wait_point.y()
        if dx * dx + dy * dy <= TASK2_CAR_ATTACH_DISTANCE_MM ** 2:
            self._uav_display_mode = "on_car"
            self._sync_uav_to_car()

    def _apply_uav_display_rules(self, frame: DcpFrame, data: dict[str, int]) -> None:
        """根据任务阶段覆盖界面坐标，不修改保存下来的原始遥测记录。"""
        phase = self._phase_from_frame(frame, data)
        position_updated = self._frame_updates_uav_position(frame, data)

        if self._task_mode == 1:
            if phase == 6:  # 伴飞：立即吸附到小车
                self._uav_display_mode = "on_car"
            elif phase == 7:  # 抛投：释放，后续按无人机实测位置显示
                self._uav_display_mode = "free"
            if self._uav_display_mode == "on_car":
                self._sync_uav_to_car()
            return

        if self._task_mode != 2:
            return

        if phase == 11:  # 从小车起飞：记录当前车上位置，开放左下方返航区域
            if self._uav_display_mode == "on_car" and "x_mm" in self.model.car and "y_mm" in self.model.car:
                self._task2_return_anchor = QPointF(
                    float(self.model.car["x_mm"]), float(self.model.car["y_mm"])
                )
            elif (
                self._task2_return_anchor is None
                and "x_mm" in self.model.uav
                and "y_mm" in self.model.uav
            ):
                self._task2_return_anchor = QPointF(
                    float(self.model.uav["x_mm"]), float(self.model.uav["y_mm"])
                )
            self._uav_display_mode = "released"
            self._task2_wait_point = None

        if self._uav_display_mode == "on_car":
            self._sync_uav_to_car()
            return

        if self._uav_display_mode == "waiting_c":
            if self._task2_wait_point is not None:
                self._set_uav_point(self._task2_wait_point)
            self._maybe_attach_task2_to_car()
            return

        if not position_updated or "x_mm" not in self.model.uav or "y_mm" not in self.model.uav:
            return

        raw_x = float(self.model.uav["x_mm"])
        raw_y = float(self.model.uav["y_mm"])
        constrained = self._task2_constrained_point(raw_x, raw_y)

        if self._uav_display_mode == "free" and self._is_near_or_beyond_task2_c(raw_x, raw_y):
            self._task2_wait_point = constrained
            self._uav_display_mode = "waiting_c"
            self._set_uav_point(constrained)
            self._maybe_attach_task2_to_car()
        elif self._uav_display_mode == "free":
            # 飞向 C 的途中不能越过 CD 右侧或 C 上方10cm；此时不限制下边界，
            # 因为任务起点 H 本来位于 D 点下方。
            approach_x = float(TASK2_C_X_MM) if raw_x > TASK2_C_X_MM + TASK2_LIMIT_MARGIN_MM else raw_x
            approach_y = float(TASK2_C_Y_MM) if raw_y > TASK2_C_Y_MM + TASK2_LIMIT_MARGIN_MM else raw_y
            self._set_uav_point(QPointF(approach_x, approach_y))
        elif self._uav_display_mode == "released":
            self._set_uav_point(self._task2_return_point(raw_x, raw_y))

    def _start_path_playback(self) -> None:
        """延时结束，开始开环播放小车轨迹。"""
        if not self._path_points:
            return
        self._path_index = 0
        self._path_playback_started_at = time.monotonic()
        # 延时结束时立即把车放到A点；之后按绝对时间轴推进，避免定时器误差累计。
        self._play_path_tick()
        if self._path_index < len(self._path_points):
            self._path_timer.start(20)

    @staticmethod
    def _nearest_path_index(points: list[QPointF], target: QPointF, begin: int) -> int:
        if begin >= len(points):
            return len(points) - 1
        return min(
            range(begin, len(points)),
            key=lambda index: (points[index] - target).manhattanLength(),
        )

    def _configure_path_timing(self, task_mode: int) -> None:
        """保留Excel空间轨迹，只按实测分段时间重新分配相邻点播放间隔。"""
        step_count = max(0, len(self._path_points) - 1)
        self._path_delay_ms = PATH_PLAYBACK_DELAY_MS
        self._path_step_intervals = [PATH_PLAYBACK_INTERVAL_MS] * step_count
        self._path_cumulative_ms = [index * PATH_PLAYBACK_INTERVAL_MS for index in range(len(self._path_points))]
        self._path_timing_checkpoints = []
        if task_mode != 2 or step_count == 0:
            return

        self._path_delay_ms = TASK2_START_TO_A_MS
        segment_start = 0
        for name, target, duration_ms in TASK2_SEGMENT_TIMING:
            segment_end = self._nearest_path_index(self._path_points, target, segment_start + 1)
            transitions = segment_end - segment_start
            if transitions <= 0:
                continue
            # 用累计四舍五入分配整数毫秒，使每段所有步长之和严格等于实测时长。
            for offset in range(transitions):
                begin_ms = round(offset * duration_ms / transitions)
                end_ms = round((offset + 1) * duration_ms / transitions)
                self._path_step_intervals[segment_start + offset] = max(1, end_ms - begin_ms)
            self._path_timing_checkpoints.append((name, segment_end, duration_ms))
            segment_start = segment_end
        cumulative = 0
        self._path_cumulative_ms = [0]
        for interval in self._path_step_intervals:
            cumulative += interval
            self._path_cumulative_ms.append(cumulative)

    def _play_path_tick(self) -> None:
        """按绝对时间更新小车位置；刷新延迟不会累计到后续路径。"""
        if self._path_index >= len(self._path_points):
            self._path_timer.stop()
            return
        if self._path_playback_started_at is None or not self._path_cumulative_ms:
            target_index = self._path_index
        else:
            elapsed_ms = max(0, round((time.monotonic() - self._path_playback_started_at) * 1000))
            target_index = min(
                len(self._path_points) - 1,
                max(0, bisect_right(self._path_cumulative_ms, elapsed_ms) - 1),
            )
            if target_index < self._path_index:
                return

        # 正常情况每次只推进一个点；若界面曾卡顿，则补处理经过的点并直接追上绝对时间。
        while self._path_index <= target_index:
            pt = self._path_points[self._path_index]
            self.model.car["x_mm"] = int(pt.x())
            self.model.car["y_mm"] = int(pt.y())
            if self._task_mode == 1 and self._uav_display_mode == "on_car":
                self._sync_uav_to_car()
            elif self._task_mode == 2:
                self._maybe_attach_task2_to_car()
                if self._uav_display_mode == "on_car":
                    self._sync_uav_to_car()
            self._path_index += 1

        if self._path_index >= len(self._path_points):
            self._path_timer.stop()
            self._path_playback_started_at = None

    def _on_serial_error(self, error: QSerialPort.SerialPortError) -> None:
        if error != QSerialPort.SerialPortError.NoError and self.port.isOpen():
            self.statusBar().showMessage(f"串口错误：{self.port.errorString()}")

    def _reset_session(self) -> None:
        self.replay_timer.stop()
        self._path_timer.stop()
        self._delay_timer.stop()
        self._reset_uav_display_rules()
        self._last_car_start_key = None
        self._path_points = []
        self._path_index = 0
        self._path_delay_ms = PATH_PLAYBACK_DELAY_MS
        self._path_step_intervals = []
        self._path_cumulative_ms = []
        self._path_timing_checkpoints = []
        self._path_playback_started_at = None
        self.task_label.setText("等待接收...")
        self.model.reset()
        self.logger.close()
        self.arena.update()

    def _load_replay(self) -> None:
        path, _ = QFileDialog.getOpenFileName(
            self, "选择地面站 JSONL 遥测日志", str(Path(__file__).resolve().parent / "logs"), "JSON Lines (*.jsonl)"
        )
        if not path:
            return
        try:
            with Path(path).open("r", encoding="utf-8") as stream:
                records = [json.loads(line) for line in stream if line.strip()]
            required = {"message_type", "source", "session_id", "seq", "sender_ms", "flags", "raw_hex"}
            if not records or any(not required.issubset(item) for item in records):
                raise ValueError("文件不包含本程序写入的 DCP JSONL 记录")
        except (OSError, ValueError, json.JSONDecodeError) as exc:
            QMessageBox.warning(self, "无法回放", str(exc))
            return
        if self.port.isOpen():
            self.port.close()
            self.connect_button.setText("连接 DL-20")
        self.model.reset()
        self.logger.close()
        self.replay_records, self.replay_index = records, 0
        self.connection_label.setText("▶ 正在回放日志")
        self.connection_label.setStyleSheet("color:#4267b2; font-weight:bold;")
        self._play_next_record()

    @staticmethod
    def _replay_delay_ms(current: dict, following: dict) -> int:
        """Preserve recorded pacing, while keeping malformed logs responsive."""
        try:
            begin = datetime.fromisoformat(current["local_time"])
            end = datetime.fromisoformat(following["local_time"])
            return max(15, min(2000, round((end - begin).total_seconds() * 1000)))
        except (KeyError, TypeError, ValueError):
            return 35

    def _play_next_record(self) -> None:
        if self.replay_index >= len(self.replay_records):
            self.replay_timer.stop()
            self.connection_label.setText("▶ 回放结束")
            self.connection_label.setStyleSheet("color:#4267b2; font-weight:bold;")
            return
        item = self.replay_records[self.replay_index]
        self.replay_index += 1
        frame = DcpFrame(
            message_type=int(item["message_type"]), flags=int(item["flags"]), source=int(item["source"]),
            destination=0, session_id=int(item["session_id"]), seq=int(item["seq"]),
            sender_ms=int(item["sender_ms"]), payload=bytes.fromhex(item["raw_hex"])[18:-3],
            raw=bytes.fromhex(item["raw_hex"]),
        )
        self._handle_frame(frame, save=False)
        if self.replay_index < len(self.replay_records):
            self.replay_timer.start(self._replay_delay_ms(item, self.replay_records[self.replay_index]))

    def _refresh_display(self) -> None:
        now = time.monotonic()
        latest = max(self.model.last_rx.values(), default=0.0)
        if self.port.isOpen() and latest and now - latest > 2.0:
            self.connection_label.setText("● 串口已连，遥测超时")
            self.connection_label.setStyleSheet("color:#c78411; font-weight:bold;")
        elif self.port.isOpen():
            self.connection_label.setText(
                f"● 已连接 {self.port.portName()} · RX {self.serial_rx_bytes}B/{self.serial_rx_frames}帧"
            )
            self.connection_label.setStyleSheet("color:#138a4a; font-weight:bold;")
        phase = self.model.uav.get("phase")
        status_text = UAV_STATUS_LABELS.get(phase, "待机") if phase is not None else "待机"
        self._show_uav_status(status_text)
        self.arena.update()

    def _clear_status_log(self) -> None:
        self.status_log.clear()

    def closeEvent(self, event) -> None:  # noqa: N802 - Qt override
        self.logger.close()
        if self._speech is not None:
            self._speech.stop()
            self._speech_queue.clear()
        if self.port.isOpen():
            self.port.close()
        event.accept()


if __name__ == "__main__":
    app = QApplication(sys.argv)
    app.setStyle("Fusion")
    window = MainWindow()
    if sys.platform == "win32":
        window.show()
    else:
        window.showFullScreen()
    sys.exit(app.exec())
